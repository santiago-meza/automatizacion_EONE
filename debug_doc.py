from docx import Document

DOC = "templates/positivo_cnv.docx"


doc = Document(DOC)


print("\n====================")
print("NORMAL PARAGRAPHS")
print("====================")

for i, p in enumerate(doc.paragraphs):
    print(f"\nParagraph {i}:")
    print(repr(p.text))

    print("RUNS:")
    for j, r in enumerate(p.runs):
        print(j, repr(r.text))


print("\n====================")
print("TABLES")
print("====================")

for ti, table in enumerate(doc.tables):

    print(f"\nTABLE {ti}")

    for ri, row in enumerate(table.rows):

        for ci, cell in enumerate(row.cells):

            for pi, p in enumerate(cell.paragraphs):

                print(
                    f"Cell {ri},{ci} paragraph {pi}:",
                    repr(p.text)
                )

                for j, r in enumerate(p.runs):
                    print("   RUN", j, repr(r.text))