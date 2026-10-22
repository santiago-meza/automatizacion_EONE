from docx import Document
from docx.text.paragraph import Paragraph
from copy import deepcopy


def insert_table_at_placeholder(report_path, table_path, output_path):

    report = Document(report_path)
    table_doc = Document(table_path)

    placeholder = "{{ CNV_TABLE }}"

    inserted = False


    for paragraph in report.paragraphs:

        full_text = "".join(
            run.text
            for run in paragraph.runs
        )

        if placeholder in full_text:

            print("CNV_TABLE placeholder found")


            # Remove placeholder text
            for run in paragraph.runs:
                run.text = ""


            body = report._body._body

            index = list(body).index(
                paragraph._p
            )


            for element in table_doc.element.body:

                if element.tag.endswith("tbl"):

                    body.insert(
                        index + 1,
                        deepcopy(element)
                    )

                    index += 1


            inserted = True
            break


    if not inserted:
        raise ValueError(
            "Placeholder {{ CNV_TABLE }} not found in body."
        )


    report.save(output_path)

# Títulos de los apartados del informe. Cada apartado va del título hasta
# el próximo título (o el final del documento).
SECTION_TITLES = {
    "Prueba solicitada",
    "Finalidad del estudio",
    "Método",
    "Resultado",
    "Interpretación",
    "Referencias",
    "Nota",
}


def keep_sections_together(report_path):
    """
    Evita que un apartado (Interpretación, Referencias, Nota, etc.) quede
    cortado entre dos hojas: si no entra entero en la hoja, Word lo pasa
    completo a la siguiente.

    Se marca cada párrafo del apartado con "Conservar líneas juntas" y
    todos menos el último con "Conservar con el siguiente". Una tabla
    corta la cadena (la tabla de CNVs puede ocupar más de una hoja).
    Si un apartado es más largo que una hoja, Word lo corta igual.
    """

    report = Document(report_path)

    chain = []

    def close_chain():
        # los párrafos vacíos del final no se encadenan
        while chain and not chain[-1].text.strip():
            chain.pop()

        for i, paragraph in enumerate(chain):
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = i < len(chain) - 1

        chain.clear()

    for element in report.element.body.iterchildren():

        if element.tag.endswith("}tbl"):
            close_chain()
            continue

        if not element.tag.endswith("}p"):
            continue

        paragraph = Paragraph(element, report)

        if paragraph.text.strip() in SECTION_TITLES:
            close_chain()
            chain.append(paragraph)

        elif chain:
            chain.append(paragraph)

    close_chain()

    report.save(report_path)
